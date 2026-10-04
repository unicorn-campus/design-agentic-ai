"""완료조건 '소스와 문서 일치' 점검 — 슬라이드 속 발췌 줄이 모두 현재 소스 파일에 그대로 있는지 대조함."""
import json
import sys
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[2] / "hybrid-ai-lab"
snippets = json.loads(Path(__file__).with_name("snippets.json").read_text(encoding="utf-8"))
deck = Presentation(sys.argv[1])
slide_text = "\n".join(sh.text_frame.text for s in deck.slides for sh in s.shapes if sh.has_text_frame)
bad = 0
for name, item in snippets.items():
    source = (ROOT / item["path"]).read_text(encoding="utf-8")
    for line in item["lines"]:
        core = line.strip()
        if core.startswith("# …(생략)"):
            continue
        if core not in source:
            bad += 1; print(f"[소스와 다름] {name}: {core[:80]}")
        if line not in slide_text:
            bad += 1; print(f"[슬라이드에 없음] {name}: {core[:80]}")
print(f"OK — 발췌 {len(snippets)}개의 모든 줄이 소스 · 슬라이드와 일치" if not bad else f"불일치 {bad}건")
