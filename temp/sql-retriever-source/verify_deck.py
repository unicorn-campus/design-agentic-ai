"""완료조건 '소스와 문서 일치' 점검.

1) 슬라이드 속 발췌 줄이 모두 현재 소스 파일에 그대로 있고 슬라이드에도 있는지 대조함
2) 해설(표, 카드)에 옮겨 적은 값이 소스에 그대로 있는지 대조함(FACTS)
3) 본문과 발표자 노트에 가운뎃점(·)이 남았는지 찾음(코드 발췌는 제외)

python verify_deck.py {덱}.pptx
"""
import json
import sys
from pathlib import Path

from pptx import Presentation

ROOT = Path.home() / "workspace" / "crc" / "sql-retriever"
HERE = Path(__file__).parent

# (소스 파일, 소스에 있어야 하는 글자) — 해설에 손으로 옮긴 값의 근거
FACTS = [
    ("app/application/models.py", s) for s in [
        '"INVALID_REQUEST": ErrorSpec(400, False', '"MEMBER_NOT_FOUND": ErrorSpec(404, False',
        '"LLM_UNAVAILABLE": ErrorSpec(503, True', '422, False, "질문을 처리할', '"DB_TIMEOUT": ErrorSpec(504, True',
        '"DB_ERROR": ErrorSpec(500, True', '"TIME_BUDGET_EXCEEDED": ErrorSpec(504, True',
        '"INTERNAL_ERROR": ErrorSpec(500, True']
] + [
    ("app/domain/fixed_queries.py", s) for s in [
        '("FQ-01", _any("가입일", "연령"))', '("FQ-02", _any("보유카드", "카드목록", "연회비"))',
        '"카드별" in text and _any("월", "사용")', '("FQ-04", _any("월별", "월사용", "지난달", "이번달"))',
        '("FQ-05", _any("업종", "가맹점", "어디서"))', '_HISTORY_WORDS = ("이력", "내역", "12개월")',
        '("FQ-08", _any("일별", "날짜별", "90일"))']
] + [
    ("app/infrastructure/groq_planner.py", s) for s in [
        "TEMPERATURE = 0.2", 'REASONING_EFFORT = "low"', "MAX_COMPLETION_TOKENS = 4096", "SEED = 42",
        "TIMEOUT_SECONDS = 5.0", "MAX_RETRIES = 0", '"include_reasoning": False', 'method="json_schema", strict=True']
] + [
    ("app/infrastructure/settings.py", 'GROQ_BASE_URL = "https://api.groq.com/openai/v1"'),
    ("app/infrastructure/settings.py", 'db_user: str = "sql_retriever_user"'),
    ("app/infrastructure/settings.py", 'values.get("RDB_NAME", "cardlab")'),
    ("app/infrastructure/settings.py", 'values.get("RDB_PORT", "5432")'),
    ("app/infrastructure/settings.py", 'values.get("RDB_HOST", "localhost")'),
    ("app/domain/time_budget.py", "START_RATIO = 0.05"),
    ("app/domain/time_budget.py", "TOTAL_BUDGET_SECONDS = 15.0"),
    ("app/domain/time_budget.py", "FINALIZE_SECONDS = 1.0"),
    ("app/domain/catalog.py", "ROW_LIMIT = 200"),
    ("app/infrastructure/postgres_repository.py", "POOL_MIN_SIZE = 1"),
    ("app/infrastructure/postgres_repository.py", "POOL_MAX_SIZE = 4"),
    ("app/infrastructure/postgres_repository.py", "max(1, timeout_ms)"),
    ("app/infrastructure/graph.py", "RECURSION_LIMIT = 12"),
    ("app/presentation/api.py", 'SEARCH_PATH = "/v1/sql-retriever/search"'),
    ("app/presentation/cli.py", 'f"cli-{uuid.uuid4().hex[:12]}"'),
    ("serve_api.py", 'parser.add_argument("--port", type=int, default=8000)'),
    ("app/domain/periods.py", "월 집계 2026-03-01 ~ 2026-08-31, 일별 2026-06-03 ~ 2026-08-31, 연체 이력 2025-09-01 ~ 2026-08-31"),
    ("app/domain/time_budget.py", "기존 사용 9.6초, 차감 5초 → 남은 시간 4.4초, 남은 시간 예산 -0.6초"),
    ("app/application/steps.py", 'REASON_RETRY_EXHAUSTED = "다시 계획해도 실행할 수 있는 계획을 만들지 못함"'),
    ("app/application/steps.py", "MAX_ATTEMPTS = 2"),
] + [
    ("app/application/steps.py", f'unfit("{code}")') for code in
    ["PLAN_FIELD_MISSING", "UNKNOWN_QUERY_ID", "EMPTY_SQL", "VIEW_NOT_ALLOWED", "INVALID_ROUTE", "INVALID_PERIOD"]
] + [
    ("app/infrastructure/sql_validator.py", f'"{code}"') for code in
    ["PARSE_ERROR", "EMPTY_SQL", "COMMENT_NOT_ALLOWED", "MULTIPLE_STATEMENTS", "WRITE_NOT_ALLOWED", "NOT_SELECT",
     "SYNTAX_NOT_ALLOWED", "FOR_UPDATE_NOT_ALLOWED", "SELECT_INTO_NOT_ALLOWED", "FORBIDDEN_FUNCTION",
     "FUNCTION_NOT_ALLOWED", "SCHEMA_NOT_ALLOWED", "VIEW_NOT_ALLOWED", "COLUMN_NOT_ALLOWED"]
] + [
    ("app/infrastructure/groq_planner.py", f'PlannerFailure("{kind}"') for kind in ["TIMEOUT", "CONNECTION", "OTHER"]
] + [
    ("app/infrastructure/groq_planner.py", s) for s in ['"RATE_LIMIT", status_code=429', 'PlannerFailure("BAD_REQUEST"',
                                                        '"SERVER_ERROR" if error.status_code >= 500']
]


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    snippets = json.loads((HERE / "snippets.json").read_text(encoding="utf-8"))
    deck = Presentation(sys.argv[1])
    slide_text = "\n".join(sh.text_frame.text for s in deck.slides for sh in s.shapes if sh.has_text_frame)
    bad = 0
    code_lines = set()
    for name, item in snippets.items():
        source = (ROOT / item["path"]).read_text(encoding="utf-8")
        for line in item["lines"]:
            code_lines.add(line)
            core = line.strip()
            if core.startswith("# …(생략)"):
                continue
            if core not in source:
                bad += 1
                print(f"[소스와 다름] {name}: {core[:80]}")
            if line not in slide_text:
                bad += 1
                print(f"[슬라이드에 없음] {name}: {core[:80]}")
    for path, needle in FACTS:
        if needle not in (ROOT / path).read_text(encoding="utf-8"):
            bad += 1
            print(f"[해설 근거 없음] {path}: {needle}")
    dots = 0
    for i, s in enumerate(deck.slides, 1):
        texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
        if s.has_notes_slide:
            texts.append(s.notes_slide.notes_text_frame.text)
        for t in texts:
            for line in t.splitlines():
                if "·" in line and line not in code_lines:
                    dots += 1
                    print(f"[가운뎃점] {i}쪽: {line[:80]}")
    print(f"발췌 {len(snippets)}개, 해설 근거 {len(FACTS)}개 대조 → 불일치 {bad}건, 가운뎃점 {dots}건")
    sys.exit(1 if bad or dots else 0)


if __name__ == "__main__":
    main()
