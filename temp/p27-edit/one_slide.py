"""교재에서 한 쪽만 남긴 사본을 만듦(렌더링 확인용). 실행: python one_slide.py 원본.pptx 쪽번호 출력.pptx"""
import sys
from pptx import Presentation
src, no, out = sys.argv[1], int(sys.argv[2]), sys.argv[3]
p = Presentation(src)
ids = p.slides._sldIdLst
for i, sid in reversed(list(enumerate(list(ids)))):
    if i != no - 1:
        p.part.drop_rel(sid.rId)
        ids.remove(sid)
p.save(out)
print("저장", out, len(p.slides), "쪽")
