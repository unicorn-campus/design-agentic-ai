"""교재 27쪽(스키마 도출 프롬프트 흐름) 아래쪽을 바꿈.

- 지움: '사람이 확정하는 곳 6번' 목록, 'AI가 지키는 규칙' 상자(옮겨 다시 만듦), '승인 화면 예(Mermaid 그림)'
- 넣음: 왼쪽 '단계별 작업' 표(9단계마다 AI가 하는 일, 사람이 판단할 일), 오른쪽 'AI가 지키는 규칙'
- 기존 도형을 복제해 글꼴(Pretendard), 크기, 색을 그대로 씀
실행: python edit_p27.py 원본.pptx 출력.pptx
"""

import copy
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

SRC, OUT = sys.argv[1], sys.argv[2]
SLIDE_NO = int(sys.argv[3]) if len(sys.argv) > 3 else 27  # 한 쪽 사본으로 시험할 때는 1

STEPS = [
    ("1. 입력 받기", "청크, 질문 출처, 키를 찾아 추천하고 개인정보 점검", "추천 입력값 확인"),
    ("2. 역량 질문", "사용자 질문을 틀로 묶어 범위별 질문 초안 작성", "질문과 범위 확정"),
    ("3. 단위와 표본", "추출 단위 규칙, 묶음, 발견용과 시험용 표본 제안", "단위, 묶음, 표본 승인"),
    ("4. 후보 뽑기", "발견용 표본을 LLM에 넣어 노드, 관계 후보 수집", None),
    ("5. 스키마 초안", "공통 개념을 먼저 정하고 병합, 제거, 참조, 추가 추천", "초안과 제약 확정"),
    ("6. 기계 검증", "스키마 꼴, 이름 규칙, 공통 노드 일치, 그림 확인", None),
    ("7. 기대 답", "시험용 표본마다 질문별 정답을 원문 그대로 초안", "기대 답, 합격 기준 확정"),
    ("8. 채우기 시험", "초안으로 추출해 질문별 점수와 연결 검사", "통과 판정 또는 설명 고치기"),
    ("9. 저장", "스키마 파일 저장 후 다시 읽어 비교, 승인 기록", None),
]
RULES = [
    "후보 뽑기에는 역량 질문만 넘김",
    "판단은 빈도 아닌 역량 질문 기준",
    "시험은 안 쓴 표본, 답 없는 문서 포함",
    "청크는 표본만 추출 단위로 나눔",
    "키는 .env에서 읽기만 함",
    "승인은 검토 파일로 받음",
    "Neo4j 판이 지원하는 제약만 씀",
]
NAVY, BLUE, TEXT, GRAY = RGBColor(0x1E, 0x2A, 0x5C), RGBColor(0x2E, 0x74, 0xC6), RGBColor(0x2B, 0x32, 0x42), RGBColor(0x7A, 0x85, 0x99)

# 배치(인치)
TOP, HEAD_H, ROW_H, GAP = 4.22, 0.40, 0.26, 0.015
X0, COL = 0.56, [1.70, 5.30, 2.95]  # 단계, AI가 하는 일, 사람이 판단할 일
PAD = 0.08
TABLE_W = sum(COL) + PAD * 2
RX, RW = 10.81, 4.65  # 오른쪽 규칙 상자
ROWS_TOP = TOP + HEAD_H + 0.06
ROWS_END = ROWS_TOP + len(STEPS) * (ROW_H + GAP) - GAP
SHIFT = 0.10  # 결과 상자와 각주를 아래로 내리는 거리


def clone(slide, shape, x, y, w, h):
    el = copy.deepcopy(shape._element)
    slide.shapes._spTree.append(el)
    new = slide.shapes[-1]
    new.left, new.top, new.width, new.height = Inches(x), Inches(y), Inches(w), Inches(h)
    return new


def set_text(shape, text, color=None, bold=None):
    tf = shape.text_frame
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    p = tf.paragraphs[0]
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)
    r = p.runs[0]
    r.text = text
    if color is not None:
        r.font.color.rgb = color
    if bold is not None:
        r.font.bold = bold


def set_lines(shape, lines):
    tf = shape.text_frame
    first = tf.paragraphs[0]
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    for r in first.runs[1:]:
        r._r.getparent().remove(r._r)
    first.runs[0].text = "• " + lines[0]
    for line in lines[1:]:
        el = copy.deepcopy(first._p)
        tf._txBody.append(el)
        tf.paragraphs[-1].runs[0].text = "• " + line


def main() -> None:
    prs = Presentation(SRC)
    s = prs.slides[SLIDE_NO - 1]
    by = {sh.name: sh for sh in s.shapes}
    assert by["Text 57"].text_frame.text == "사람이 확정하는 곳 6번", "27쪽 구성이 예상과 다름"
    tpl = {k: by[k] for k in ("Shape 56", "Text 57", "Shape 58", "Text 61", "Shape 82", "Text 83", "Shape 84", "Text 85")}
    # 1) 아래쪽 세 칸(Shape 56 ~ Text 100)을 지움. 템플릿은 먼저 복제한 뒤 지움
    names = [f"Shape {i}" for i in range(56, 101)] + [f"Text {i}" for i in range(56, 101)]
    kept = {k: copy.deepcopy(v._element) for k, v in tpl.items()}
    for n in names:
        if n in by:
            by[n]._element.getparent().remove(by[n]._element)
    for k, el in kept.items():  # 복제용으로 템플릿을 다시 넣어 두고 마지막에 지움
        s.shapes._spTree.append(el)
    t = {sh.name: sh for sh in s.shapes if sh.name in kept}
    # 2) 표 머리줄: 단계 | AI가 하는 일(남색) | 사람이 판단할 일(파랑)
    x = X0
    for (label, w), (box, txt) in zip((("단계", COL[0] + PAD), ("AI가 하는 일", COL[1]), ("사람이 판단할 일", COL[2] + PAD)),
                                      (("Shape 82", "Text 83"), ("Shape 82", "Text 83"), ("Shape 56", "Text 57"))):
        clone(s, t[box], x, TOP, w - 0.03, HEAD_H)
        set_text(clone(s, t[txt], x, TOP, w - 0.03, HEAD_H), label)
        x += w
    # 3) 단계 줄
    for i, (step, ai, human) in enumerate(STEPS):
        y = ROWS_TOP + i * (ROW_H + GAP)
        clone(s, t["Shape 58"], X0, y, TABLE_W, ROW_H)
        set_text(clone(s, t["Text 61"], X0 + PAD, y, COL[0], ROW_H), step, NAVY, True)
        set_text(clone(s, t["Text 61"], X0 + PAD + COL[0], y, COL[1], ROW_H), ai, TEXT, False)
        set_text(clone(s, t["Text 61"], X0 + PAD + COL[0] + COL[1], y, COL[2], ROW_H),
                 human or "(AI 단독)", BLUE if human else GRAY, bool(human))
    # 4) 오른쪽 규칙 상자
    clone(s, t["Shape 82"], RX, TOP, RW, HEAD_H)
    set_text(clone(s, t["Text 83"], RX, TOP, RW, HEAD_H), "AI가 지키는 규칙")
    clone(s, t["Shape 84"], RX, ROWS_TOP, RW, ROWS_END - ROWS_TOP)
    set_lines(clone(s, t["Text 85"], RX + 0.18, ROWS_TOP + 0.05, RW - 0.30, ROWS_END - ROWS_TOP - 0.10), RULES)
    # 템플릿 지움
    for sh in list(s.shapes):
        if sh._element in kept.values():
            sh._element.getparent().remove(sh._element)
    # 5) 결과 상자와 각주를 조금 내림
    for n in ("Shape 101", "Shape 102", "Text 103", "Text 104"):
        sh = next(x for x in s.shapes if x.name == n)
        sh.top = sh.top + Inches(SHIFT)
    prs.save(OUT)
    print("저장", OUT, "| 표 끝", round(ROWS_END, 2), "in, 결과 상자 시작", round(7.12 + SHIFT, 2), "in")


if __name__ == "__main__":
    main()
