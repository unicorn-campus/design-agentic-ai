"""빌드 후처리: 개체 틀 번호 정리(가이드 6-14) · zip 빈 폴더 항목 제거 · adj 범위 검사 · 붙여넣기 호환 점검(6-13)."""
import re
import shutil
import sys
import zipfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Emu, Inches, Pt

src = sys.argv[1]
tmp = src + ".tmp"
bad_adj = []
with zipfile.ZipFile(src) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        if item.filename.endswith("/"):
            continue  # 빈 폴더 항목은 PowerPoint가 손상으로 볼 수 있어 뺌
        data = zin.read(item.filename)
        if re.match(r"ppt/(slides|slideLayouts)/slide(Layout)?\d+\.xml$", item.filename):
            x = data.decode("utf8")
            x = re.sub(r'<p:ph\s+idx="\d+"\s+type="title"[^>]*/>', '<p:ph type="title"/>', x)
            x = re.sub(r'<p:ph type="sldNum" sz="quarter" idx="\d+"/>', '<p:ph type="sldNum" idx="4"/>', x)
            for v in re.findall(r'<a:gd name="adj" fmla="val (\d+)"', x):
                if int(v) > 50000:
                    bad_adj.append((item.filename, v))
            data = x.encode("utf8")
        zout.writestr(item, data)
shutil.move(tmp, src)

p = Presentation(src)


def fix_cover(slide) -> bool:
    """제목 장표(맨 앞 도형이 화면 전체 사각형)의 제목 개체 틀을 가운데 큰 흰 글자로 옮기고 쪽 번호를 지움.

    pptxgenjs는 개체 틀에 준 위치 · 색을 무시하고 레이아웃 값을 써서, 빌드 뒤에 직접 고침.
    """

    first = slide.shapes[0] if len(slide.shapes) else None
    if first is None or first.width < Emu(14000000) or first.height < Emu(8000000):
        return False
    title = slide.shapes.title
    title.left, title.top, title.width, title.height = Inches(1.0), Inches(3.45), Inches(13.5), Inches(1.2)
    for para in title.text_frame.paragraphs:
        for run in para.runs:
            run.font.size = Pt(44)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for shape in list(slide.placeholders):
        if shape.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
            shape._element.getparent().remove(shape._element)
    return True


covers = [i for i, s in enumerate(p.slides, 1) if fix_cover(s)]
if covers:
    p.save(src)
    p = Presentation(src)
err = []
if (p.slide_width, p.slide_height) != (14630400, 8229600):
    err.append("슬라이드 크기가 16x9인치가 아님")
for i, s in enumerate(p.slides, 1):
    if s.slide_layout.name != "내용_단쪽":
        err.append(f"{i}쪽 레이아웃 이름: {s.slide_layout.name}")
    if s.shapes.title is None or not s.shapes.title.text.strip():
        err.append(f"{i}쪽 제목 개체 틀 비었음")
xml = "".join(s._element.xml for s in p.slides)
if "schemeClr" in xml:
    err.append("테마 색(schemeClr) 사용")
for f, v in bad_adj:
    err.append(f"adj 범위 초과 {f}: {v}")
print("\n".join(err) or f"OK — {len(p.slides)}쪽, 개체 틀 정리 · 호환 점검 통과")
